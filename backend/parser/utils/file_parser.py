# parser/utils/file_parser.py

import PyPDF2
import docx
import io


class DocumentParser:
    """Parse PDF and DOCX files to extract text"""
    
    @staticmethod
    def parse_pdf(file_content):
        """
        Extract text from PDF file
        
        :param file_content: File bytes or file object
        :return: Extracted text as string
        """
        try:
            # Create a PDF reader object
            if isinstance(file_content, bytes):
                file_content = io.BytesIO(file_content)
            
            pdf_reader = PyPDF2.PdfReader(file_content)
            
            # Extract text from all pages
            text = ""
            for page_num in range(len(pdf_reader.pages)):
                page = pdf_reader.pages[page_num]
                text += page.extract_text() + "\n"
            
            return text.strip()
        
        except Exception as e:
            raise ValueError(f"Error parsing PDF: {str(e)}")
    
    @staticmethod
    def parse_docx(file_content):
        """
        Extract text from DOCX file
        
        :param file_content: File bytes or file object
        :return: Extracted text as string
        """
        try:
            # Create a Document object
            if isinstance(file_content, bytes):
                file_content = io.BytesIO(file_content)
            
            doc = docx.Document(file_content)
            
            # Extract text from all paragraphs
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + " "
                text += "\n"
            
            return text.strip()
        
        except Exception as e:
            raise ValueError(f"Error parsing DOCX: {str(e)}")
    
    @staticmethod
    def parse_document(file_obj, filename):
        """
        Parse document based on file extension
        
        :param file_obj: Django UploadedFile object or file content
        :param filename: Name of the file
        :return: Tuple of (extracted_text, file_type)
        """
        # Get file extension
        file_extension = filename.lower().split('.')[-1]
        
        # Read file content
        if hasattr(file_obj, 'read'):
            file_content = file_obj.read()
        else:
            file_content = file_obj
        
        # Parse based on extension
        if file_extension == 'pdf':
            text = DocumentParser.parse_pdf(file_content)
            return text, 'pdf'
        
        elif file_extension in ['docx', 'doc']:
            text = DocumentParser.parse_docx(file_content)
            return text, 'docx'
        
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")